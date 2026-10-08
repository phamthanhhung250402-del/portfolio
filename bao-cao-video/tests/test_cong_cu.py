"""Kiểm thử: không đè cột công thức, ghép ID đúng, mix cộng đủ tổng số video.

Chạy:  python3 -m unittest discover -s tests -v   (trong thư mục bao-cao-video)
"""
from __future__ import annotations

import random
import shutil
import sys
import tempfile
import unittest
import warnings
from pathlib import Path

import openpyxl

GOC = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(GOC))
sys.path.insert(0, str(GOC / "tests"))

import tao_du_lieu_gia  # noqa: E402
from baocao import chay as chay_mod  # noqa: E402
from baocao.config import LoiCauHinh, doc_brand  # noqa: E402
from baocao.exports import doc_export  # noqa: E402
from baocao.ids import fb_ids, tiktok_id  # noqa: E402
from baocao.mix import UngVien, phan_bo  # noqa: E402
from baocao.util import doc_so, nhan_dien_dong_thang, so_tsv  # noqa: E402


def _o_cong_thuc(ws):
    return {(c.row, c.column): c.value for row in ws.iter_rows() for c in row
            if c.data_type == "f" or (isinstance(c.value, str) and c.value.startswith("="))}


class LuongDayDu(unittest.TestCase):
    """Chạy toàn bộ luồng trên dữ liệu giả tạo mới trong thư mục tạm."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp(prefix="baocao-test-"))
        cls.gia = tao_du_lieu_gia.tao(cls.tmp / "input")
        chay_mod.in_ra = lambda *a: None
        cls.kq = chay_mod.chay("topzone", mo_report=False, thu_muc_input=cls.tmp / "input",
                               thu_muc_output=cls.tmp / "output")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    # ---------------------------------------------------------- công thức
    def test_khong_xuat_cot_cong_thuc(self):
        lich, khoi = self.kq["lich"], self.kq["khoi"]
        self.assertTrue(khoi, "phải có ít nhất 1 khối TSV")
        for k in khoi:
            for r in range(k.dong_dau, k.dong_cuoi + 1):
                for c in range(k.cot_dau, k.cot_cuoi + 1):
                    self.assertFalse(lich.la_cong_thuc(r, c), f"Khối {k.so} chứa ô công thức tại dòng {r} cột {c}")

    def test_dan_tsv_khong_de_cong_thuc_va_giu_gia_tri_cu(self):
        """Mô phỏng dán từng khối vào sheet: công thức còn nguyên, dòng không có số mới giữ giá trị cũ."""
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(self.gia["sheet"])
        ws = wb["03 Content Calendars"]
        truoc = _o_cong_thuc(ws)
        cu = {(r, c): ws.cell(r, c).value for r in range(1, ws.max_row + 1) for c in range(1, ws.max_column + 1)}
        ghep = self.kq["ghep"]
        for k in self.kq["khoi"]:
            dong = k.noi_dung.rstrip("\n").split("\n")
            self.assertEqual(len(dong), k.dong_cuoi - k.dong_dau + 1)
            for i, d in enumerate(dong):
                r = k.dong_dau + i
                o = d.split("\t")
                self.assertEqual(len(o), k.cot_cuoi - k.cot_dau + 1)
                for j, x in enumerate(o):
                    c = k.cot_dau + j
                    if r not in ghep.cap_nhat:
                        self.assertEqual(x, so_tsv(cu[(r, c)], ","), f"Ô dòng {r} cột {c} bị đổi dù không có số mới")
                    ws.cell(r, c, x if x != "" else None)
        sau = _o_cong_thuc(ws)
        self.assertEqual(truoc, sau, "Có công thức bị thay đổi sau khi dán TSV")

    def test_huong_dan_dan(self):
        txt = (self.kq["out"] / "cap-nhat-tab03" / "huong-dan-dan.txt").read_text(encoding="utf-8")
        for k in self.kq["khoi"]:
            self.assertIn(f"Dán khối {k.so} vào ô {k.o_dau}", txt)
        self.assertIn("Dán khối 1 vào ô P", txt)

    # ---------------------------------------------------------- ghép ID
    def test_ghep_dung_so_lieu(self):
        ghep = self.kq["ghep"]
        loi = {id(v) for v in self.gia["loi"].values()}
        dem = 0
        for v in self.gia["videos"]:
            if v["thang"] != tao_du_lieu_gia.THANG_CAP_NHAT or v["status"] != "Đã air" or not v["nguoi"].startswith("Hưng"):
                continue
            if id(v) in loi:
                continue
            cn = ghep.cap_nhat.get(v["dong"])
            self.assertIsNotNone(cn, f"Video dòng {v['dong']} không được ghép")
            self.assertEqual(cn["tiktok_views"], v["views"])
            self.assertEqual(cn["tiktok_likes"], v["likes"])
            self.assertEqual(cn["fb_xem_3s"], v["fb_3s"])
            self.assertEqual(cn["fb_reach"], v["reach"])
            self.assertAlmostEqual(cn["fb_xem_tb"], v["xem_tb"])
            self.assertAlmostEqual(cn["chi_phi_tiktok"], v["cp_tt"], delta=1.5)  # QC tách 2 nhóm được cộng lại
            self.assertEqual(cn["chi_phi_fb"], v["cp_fb"])
            self.assertEqual(cn["fb_ads_luot_phat"], v["plays"])
            dem += 1
        self.assertGreater(dem, 8)

    def test_chi_cap_nhat_video_cua_toi(self):
        ghep = self.kq["ghep"]
        cua_nguoi_khac = {v["dong"] for v in self.gia["videos"] if not v["nguoi"].startswith("Hưng")}
        self.assertFalse(cua_nguoi_khac & set(ghep.cap_nhat), "Cập nhật nhầm dòng của người khác (vd 'Hùng')")
        self.assertTrue(any(v["nguoi"] == "Hùng" for v in self.gia["videos"]), "dữ liệu giả phải có 'Hùng' để kiểm tra")

    def test_bao_video_khong_ghep_duoc(self):
        ghep = self.kq["ghep"]
        ly_do = {v.dong: " | ".join(l) for v, l in ghep.khong_ghep_video}
        loi = self.gia["loi"]
        self.assertIn("Thiếu link TikTok", ly_do[loi["thieu_link_tt"]["dong"]])
        self.assertIn("rút gọn", ly_do[loi["link_rut_gon"]["dong"]])
        self.assertIn("chia sẻ", ly_do[loi["fb_share"]["dong"]])
        self.assertIn("Chưa có TikTok views cho ID", ly_do[loi["sai_id"]["dong"]])
        self.assertTrue(any("7999999999999999999" in m for _, m, _ in ghep.khong_ghep_dong),
                        "Dòng quảng cáo có ID lạ phải được báo")

    def test_mix_cong_du_tong(self):
        pt = self.kq["phan_tich"]
        self.assertEqual(sum(d["so_video"] for d in pt["mix"]), pt["thang_sau"]["tong"])
        tsv = (self.kq["out"] / "de-xuat-mix-thang-sau.tsv").read_text(encoding="utf-8").splitlines()
        self.assertEqual(tsv[0].split("\t"), ["Nhóm nội dung", "Số video", "Tỷ trọng", "Thời lượng mục tiêu", "Lý do", "Độ tin cậy"])
        self.assertEqual(sum(int(d.split("\t")[1]) for d in tsv[1:-1]), pt["thang_sau"]["tong"])
        for d in pt["mix"]:
            self.assertTrue(d["ly_do"], f"Nhóm {d['ten']} thiếu lý do")
        # nhóm < 5 video phải gắn nhãn thử nghiệm
        for d in pt["mix"]:
            if d["nhom"]["n"] < 5:
                self.assertEqual(d["do_tin_cay"], "Thử nghiệm - chưa đủ dữ liệu")

    def test_dau_ra_day_du(self):
        out = self.kq["out"]
        for f in ("report.html", "de-xuat-mix-thang-sau.tsv", "brief-cho-claude.md", "cap-nhat-tab03/huong-dan-dan.txt"):
            self.assertTrue((out / f).exists(), f)
        brief = (out / "brief-cho-claude.md").read_text(encoding="utf-8")
        self.assertIn("CHƯA viết kịch bản", brief)
        html = (out / "report.html").read_text(encoding="utf-8")
        self.assertIn("chart.js@4.4.1", html)


class GhepID(unittest.TestCase):
    def test_tiktok(self):
        self.assertEqual(tiktok_id("https://www.tiktok.com/@topzone.official/video/7686797298801626369?is_from_webapp=1"),
                         "7686797298801626369")
        self.assertEqual(tiktok_id("https://www.tiktok.com/@a.b/photo/7686797298801626369"), "7686797298801626369")
        self.assertIsNone(tiktok_id("https://vt.tiktok.com/ZSabc123/"))
        self.assertIsNone(tiktok_id(None))

    def test_facebook(self):
        self.assertEqual(fb_ids("https://www.facebook.com/topzone.vn/videos/1234567890123456/"), {"1234567890123456"})
        self.assertEqual(fb_ids("https://www.facebook.com/reel/1234567890123456"), {"1234567890123456"})
        self.assertEqual(fb_ids("https://www.facebook.com/100064123456789/posts/1234567890123456"), {"1234567890123456"})
        self.assertEqual(fb_ids("https://www.facebook.com/watch/?v=1234567890123456"), {"1234567890123456"})
        self.assertEqual(fb_ids("https://www.facebook.com/permalink.php?story_fbid=1234567890123456&id=100064123456789"),
                         {"1234567890123456"})
        self.assertEqual(fb_ids("100064123456789_1234567890123456"), {"1234567890123456"})
        self.assertEqual(fb_ids("https://www.facebook.com/1089295270518341"), {"1089295270518341"})
        self.assertIn("pfbid0abcdefghijklmnopqrstuvwxyz123", fb_ids("https://www.facebook.com/topzone/posts/pfbid0abcdefghijklmnopqrstuvwxyz123"))
        # Không lấy ID Trang
        self.assertNotIn("100064123456789", fb_ids("https://www.facebook.com/100064123456789/posts/1234567890123456"))


class DocSo(unittest.TestCase):
    def test_doc_so(self):
        self.assertEqual(doc_so("1.234.567", nguyen=True), 1234567)
        self.assertEqual(doc_so("407.819", nguyen=True), 407819)
        self.assertEqual(doc_so("1,234,567"), 1234567)
        self.assertEqual(doc_so("12,5"), 12.5)
        self.assertEqual(doc_so("7.4"), 7.4)
        self.assertEqual(doc_so("1.2K"), 1200)
        self.assertEqual(doc_so("3,4M"), 3_400_000)
        self.assertAlmostEqual(doc_so("45%"), 0.45)
        self.assertEqual(doc_so("9,9đ"), 9.9)
        self.assertEqual(doc_so("1234567.89", nguyen=True), 1234567.89)
        self.assertIsNone(doc_so("-"))

    def test_dong_thang(self):
        self.assertEqual(nhan_dien_dong_thang("JULY 2026 - BAU / Back to School · 25 ..."), "2026-07")
        self.assertEqual(nhan_dien_dong_thang("THÁNG 9/2026 - Launch"), "2026-09")
        self.assertIsNone(nhan_dien_dong_thang("Tuần 1"))
        self.assertIsNone(nhan_dien_dong_thang("Mayday sale"))


class NhanDienExport(unittest.TestCase):
    def test_mau_header(self):
        brand = doc_brand("topzone")
        ky_vong = {"GIA-DINH_tiktok-ads-manager.csv": "tiktok_ads", "GIA-DINH_meta-ads-manager.csv": "meta_ads",
                   "GIA-DINH_meta-business-suite.csv": "meta_business_suite"}
        for ten, loai in ky_vong.items():
            fe = doc_export(GOC / "samples" / "headers" / ten, brand.mappings)
            self.assertEqual(fe.loai, loai, ten)
            self.assertFalse(fe.thieu_truong, f"{ten} thiếu {fe.thieu_truong}")


class PhanBoMix(unittest.TestCase):
    def test_ngau_nhien_luon_cong_du(self):
        rnd = random.Random(7)
        for _ in range(3000):
            n = rnd.randint(0, 60)
            uv = []
            for i in range(rnd.randint(1, 11)):
                xh = rnd.random() < 0.7
                uv.append(UngVien(f"N{i}", diem=rnd.uniform(0, 100) if rnd.random() > 0.05 else None, xep_hang=xh,
                                  dang_chay=rnd.random() < 0.8, uu_tien_thu_nghiem=rnd.random() < 0.15,
                                  toi_thieu_them=rnd.choice([0, 0, 0, 1, 2]), so_video_lich_su=rnd.randint(0, 30)))
            kq = phan_bo(n, uv)
            self.assertEqual(sum(kq.so_video.values()), n, (n, uv))
            self.assertTrue(all(v >= 0 for v in kq.so_video.values()))

    def test_tran_va_toi_thieu(self):
        uv = [UngVien(f"N{i}", diem=d, xep_hang=True, dang_chay=True) for i, d in enumerate([95, 80, 60, 40, 20, 10])]
        uv.append(UngVien("Moi", xep_hang=False))
        kq = phan_bo(25, uv)
        self.assertEqual(sum(kq.so_video.values()), 25)
        tran = int(0.35 * 25)
        for k, v in kq.so_video.items():
            self.assertLessEqual(v, tran, k)
        for u in uv[:-1]:
            self.assertGreaterEqual(kq.so_video[u.ten], 2, u.ten)
        self.assertEqual(kq.thu_nghiem["Moi"], 5)  # 20% của 25
        self.assertGreaterEqual(kq.so_video["N0"], kq.so_video["N5"])


class TachBietThuongHieu(unittest.TestCase):
    def test_khong_dung_chung_thu_muc(self):
        tmp = Path(tempfile.mkdtemp())
        try:
            shutil.copytree(GOC / "config", tmp / "config")
            shutil.copytree(GOC / "config" / "brands" / "topzone", tmp / "config" / "brands" / "brand2")
            with self.assertRaises(LoiCauHinh):
                doc_brand("brand2", tmp)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
