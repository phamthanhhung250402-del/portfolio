"""Đọc tab 03 Content Calendars: khối tháng, dòng video, ô nào là CÔNG THỨC."""
from __future__ import annotations

import re
import unicodedata
import warnings
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from .config import Brand, LoiCauHinh, chu_sang_so, so_sang_chu
from .util import chuan_hoa, la_trong, nhan_dien_dong_thang


@dataclass
class DongVideo:
    dong: int                    # số dòng trong sheet (1-based)
    thang: str                   # '2026-09'
    gia_tri: dict                # khoá cột -> giá trị (đã tính sẵn trong file)
    cong_thuc: dict              # khoá cột -> True nếu ô là công thức

    def get(self, k, mac_dinh=None):
        v = self.gia_tri.get(k)
        return mac_dinh if la_trong(v) else v


@dataclass
class KhoiThang:
    thang: str
    tieu_de: str
    dong_tieu_de: int
    dong_dau: int                # dòng dữ liệu đầu tiên (ngay dưới dòng tiêu đề tháng)
    dong_cuoi: int               # dòng có nội dung cuối cùng của khối
    videos: list = field(default_factory=list)


@dataclass
class LichNoiDung:
    duong_dan: Path
    ten_tab: str
    cot: dict                    # khoá -> số thứ tự cột (1-based)
    khoi: list                   # [KhoiThang]
    canh_bao: list
    ws_f: object = None          # sheet đọc công thức (để kiểm tra ô công thức)
    ws_v: object = None          # sheet đọc giá trị

    def khoi_thang(self, thang: str) -> KhoiThang | None:
        for k in self.khoi:
            if k.thang == thang:
                return k
        return None

    def la_cong_thuc(self, dong: int, cot: int) -> bool:
        return _la_cong_thuc(self.ws_f.cell(dong, cot))

    def gia_tri_o(self, dong: int, cot: int):
        return self.ws_v.cell(dong, cot).value


def _la_cong_thuc(cell) -> bool:
    if cell.data_type == "f":
        return True
    v = cell.value
    return isinstance(v, str) and v.startswith("=")


def tim_tab(wb, ten: str):
    if ten in wb.sheetnames:
        return ten
    k = chuan_hoa(ten)
    for s in wb.sheetnames:
        if chuan_hoa(s) == k:
            return s
    tien_to = ten.split()[0] if ten else "03"
    for s in wb.sheetnames:
        if s.strip().startswith(tien_to):
            return s
    raise LoiCauHinh(f"Không tìm thấy tab '{ten}' trong file. Các tab hiện có: {', '.join(wb.sheetnames)}")


def doc_lich(path: Path, brand: Brand) -> LichNoiDung:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb_f = openpyxl.load_workbook(path, data_only=False)
        wb_v = openpyxl.load_workbook(path, data_only=True)
    ten_tab = tim_tab(wb_f, brand.raw.get("tab_lich", "03 Content Calendars"))
    ws_f, ws_v = wb_f[ten_tab], wb_v[ten_tab]
    dong_td = int(brand.raw.get("dong_tieu_de", 5))
    canh_bao: list[str] = []

    # --- Xác định cột theo tên tiêu đề, không thấy thì theo chữ cột mặc định
    tieu_de = {}
    for c in range(1, ws_v.max_column + 1):
        t = chuan_hoa(ws_v.cell(dong_td, c).value)
        if t and t not in tieu_de:
            tieu_de[t] = c
    cot: dict[str, int] = {}
    for k, (chu, ten) in brand.cot.items():
        c = tieu_de.get(chuan_hoa(ten))
        mac_dinh = chu_sang_so(chu)
        if c is None:
            cot[k] = mac_dinh
            thuc_te = ws_v.cell(dong_td, mac_dinh).value
            canh_bao.append(
                f"Tab 03: không thấy tiêu đề '{ten}' ở dòng {dong_td}; dùng cột {chu} "
                f"(đang ghi '{thuc_te or 'trống'}'). Nếu sai, sửa tên trong brand.yaml.")
        else:
            cot[k] = c
            if c != mac_dinh:
                canh_bao.append(f"Tab 03: cột '{ten}' đang ở {so_sang_chu(c)} (mặc định {chu}) - tool tự dùng {so_sang_chu(c)}.")

    # --- Duyệt dòng: tìm dòng tiêu đề tháng và dòng video
    khoi: list[KhoiThang] = []
    hien_tai: KhoiThang | None = None
    cot_noi_dung = [cot[k] for k in ("chu_de", "link_tiktok", "nhom_noi_dung", "status", "nguoi_lam") if k in cot]
    for r in range(dong_td + 1, ws_v.max_row + 1):
        thang = None
        for c in (1, 2, 3):
            thang = nhan_dien_dong_thang(ws_v.cell(r, c).value)
            if thang:
                tieu = str(ws_v.cell(r, c).value).strip()
                break
        if thang:
            hien_tai = KhoiThang(thang, tieu, r, r + 1, r)
            khoi.append(hien_tai)
            continue
        if hien_tai is None:
            continue
        if not any(not la_trong(ws_v.cell(r, c).value) for c in cot_noi_dung):
            continue
        hien_tai.dong_cuoi = r
        gt, ct = {}, {}
        for k, c in cot.items():
            gt[k] = ws_v.cell(r, c).value
            ct[k] = _la_cong_thuc(ws_f.cell(r, c))
        hien_tai.videos.append(DongVideo(r, hien_tai.thang, gt, ct))

    if not khoi:
        raise LoiCauHinh(
            "Không tìm thấy dòng tiêu đề tháng nào trong tab 03 (vd 'JULY 2026 - ...'). "
            "Dòng tiêu đề tháng phải nằm ở cột A, bắt đầu bằng tên tháng + năm.")
    trung = {}
    for k in khoi:
        trung.setdefault(k.thang, []).append(k)
    for t, ds in trung.items():
        if len(ds) > 1:
            canh_bao.append(f"Tab 03 có {len(ds)} khối cùng tháng {t} (dòng {', '.join(str(x.dong_tieu_de) for x in ds)}) - chỉ dùng khối đầu tiên để cập nhật.")
    return LichNoiDung(path, ten_tab, cot, khoi, canh_bao, ws_f, ws_v)


def _tu_ten(s) -> list[str]:
    # Giữ dấu tiếng Việt: "Hưng" khác "Hùng"
    return re.findall(r"\w+", unicodedata.normalize("NFC", str(s or "")).lower())


def la_cua_toi(v: DongVideo, brand: Brand) -> bool:
    """Ô Người làm có tên của bạn (vd 'Hưng', 'Hưng, Vy', 'Hưng + Vy')."""
    ten = _tu_ten(brand.raw.get("nguoi_lam", ""))
    o = _tu_ten(v.get("nguoi_lam", ""))
    n = len(ten)
    return n > 0 and any(o[i:i + n] == ten for i in range(len(o) - n + 1))


def da_air(v: DongVideo, brand: Brand) -> bool:
    return chuan_hoa(brand.raw.get("status_da_air", "Đã air")) in chuan_hoa(v.get("status", ""))
